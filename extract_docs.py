import os
from docx import Document

docs_folder = r'e:\PortfolioProjects\AI-ShopAssistant\SmartShop_AI\ReferenceDocs'
output_folder = r'e:\PortfolioProjects\AI-ShopAssistant\SmartShop_AI\ReferenceDocs'

doc_files = [
    'SmartShop_AI_Proposal_Design.docx',
    'SmartShop_AI_SRS_SDS.docx',
    'SmartShop_AI_Backend.docx',
    'SmartShop_AI_API_Documentation.docx',
    'SmartShop_AI_MLArchitecture.docx'
]

for doc_file in doc_files:
    doc_path = os.path.join(docs_folder, doc_file)
    txt_file = doc_file.replace('.docx', '.txt')
    txt_path = os.path.join(output_folder, txt_file)
    
    try:
        doc = Document(doc_path)
        with open(txt_path, 'w', encoding='utf-8') as f:
            for para in doc.paragraphs:
                if para.text.strip():
                    f.write(para.text + '\n')
            
            # Extract tables
            for table in doc.tables:
                f.write('\n[TABLE]\n')
                for row in table.rows:
                    row_text = ' | '.join([cell.text.strip() for cell in row.cells])
                    f.write(row_text + '\n')
                f.write('[/TABLE]\n\n')
        
        print(f'Extracted: {txt_file}')
    except Exception as e:
        print(f'Error extracting {doc_file}: {str(e)}')

print('Extraction complete!')
